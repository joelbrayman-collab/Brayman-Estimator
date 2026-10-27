"""Permanent document-governance contract (Joel, 27 Oct 2026).

For every document type represented by an approved master (Master Template Family V1, FG-022), that
master is mandatory: the type resolves to exactly one master, the master bytes must be the approved
ones, population changes only designated variable content, and no other writer or layout satisfies
the type. "DATA CHANGES BY PROJECT. THE APPROVED DOCUMENT FORMAT DOES NOT."
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document

from app.services import approved_masters as AM
from app.services import family_05_master

ROOT = Path(__file__).resolve().parents[1]
REGISTER = ROOT / "docs/testing/reusable-approved-document-template-family-v1-register.md"


def _register_rows():
    text = REGISTER.read_text(encoding="utf-8")
    names = re.findall(r"\| Master filename \| `MASTER DOCX/([^`]+)` \|", text)
    shas = re.findall(r"\| Master DOCX SHA-256 \| `([0-9a-f]{64})` \|", text)
    return dict(zip(names, shas))


def test_every_governed_document_type_resolves_to_its_registered_approved_master():
    register = _register_rows()
    assert len(register) == 7
    assert {m.family for m in AM.APPROVED_MASTERS.values()} == {"01", "02", "03", "04", "05", "06", "07"}
    for document_type in AM.GOVERNED_DOCUMENT_TYPES:
        master = AM.master_for(document_type)
        assert register[master.filename] == master.sha256, document_type
        assert master.filename.startswith(master.family + "_") and master.filename.endswith("_V1.docx")
    assert AM.master_for(AM.ONTARIO_CONSTRUCTION_CONTRACT).sha256 == family_05_master.FAMILY_05_MASTER_SHA256


def test_no_other_writer_or_layout_satisfies_a_governed_document_type():
    for document_type in AM.GOVERNED_DOCUMENT_TYPES:
        AM.assert_governed_route(document_type, AM.__name__)
        for writer in ("branded_issue_pdf", "app.services.proposal_pdf", "reportlab", "markdown"):
            with pytest.raises(AM.ApprovedMasterError) as raised:
                AM.assert_governed_route(document_type, writer)
            assert raised.value.block_code == AM.BLOCK_ALTERNATE_WRITER
    with pytest.raises(AM.ApprovedMasterError) as raised:
        AM.master_for("POOL_DECK_ESTIMATE_LAYOUT")
    assert raised.value.block_code == AM.BLOCK_UNKNOWN_DOCUMENT_TYPE


def test_a_file_that_is_not_the_approved_master_is_refused(tmp_path):
    master = AM.master_for(AM.CUSTOMER_FACING_ESTIMATE)
    (tmp_path / master.filename).write_bytes(b"a look-alike document")
    with pytest.raises(AM.ApprovedMasterError) as raised:
        AM.load_master_copy(AM.CUSTOMER_FACING_ESTIMATE, tmp_path)
    assert raised.value.block_code == AM.BLOCK_MASTER_SHA_MISMATCH
    with pytest.raises(AM.ApprovedMasterError) as raised:
        AM.load_master_copy(AM.CUSTOMER_FACING_ESTIMATE, tmp_path / "empty")
    assert raised.value.block_code == AM.BLOCK_MISSING_MASTER


def _synthetic_master(path: Path) -> bytes:
    document = Document()
    document.add_paragraph("CUSTOMER-FACING ESTIMATE")
    info = document.add_table(rows=2, cols=4)
    for cell, text in zip(info.rows[0].cells + info.rows[1].cells, ["Client", "CLIENT", "Project", "PROJECT", "Site", "SITE", "Date", "DATE"]):
        cell.text = text
    document.add_paragraph("Total including HST: AMOUNT")
    table = document.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Included Scope", "Description"
    for row in table.rows[1:]:
        row.cells[0].text = row.cells[1].text = AM.PLACEHOLDER
    document.add_paragraph("• First standard item.")
    document.add_paragraph("• Second standard item.")
    document.sections[0].footer.paragraphs[0].text = "CLIENT | SITE | Customer-Facing Estimate | DATE"
    buffer = io.BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    path.write_bytes(data)
    return data


def test_population_changes_only_the_designated_content_of_the_master(tmp_path, monkeypatch):
    master = AM.master_for(AM.CUSTOMER_FACING_ESTIMATE)
    data = _synthetic_master(tmp_path / master.filename)
    monkeypatch.setitem(AM.APPROVED_MASTERS, AM.CUSTOMER_FACING_ESTIMATE,
                        AM.ApprovedMaster(master.family, master.document_type, master.title, master.filename,
                                          hashlib.sha256(data).hexdigest(), master.classification))
    out = AM.populate(AM.CUSTOMER_FACING_ESTIMATE, directory=tmp_path,
                      tokens={"CLIENT": "Linda Bushel (synthetic)", "PROJECT": "Rear deck", "SITE": "1 Test Rd",
                              "DATE": "2026-10-27", "AMOUNT": "$1,130.00"},
                      tables={("Included Scope", "Description"): [("A", "a"), ("B", "b"), ("C", "c")]},
                      bullets=["Only one exclusion.", "Closing note."])
    produced = Document(io.BytesIO(out))
    text = "\n".join(p.text for p in produced.paragraphs)
    assert "Total including HST: $1,130.00" in text and "• Only one exclusion." in text and text.endswith("• Closing note.")
    assert "standard item" not in text
    scope = [t for t in produced.tables if t.rows[0].cells[0].text == "Included Scope"][0]
    assert [r.cells[0].text for r in scope.rows[1:]] == ["A", "B", "C"]
    assert produced.sections[0].footer.paragraphs[0].text == "Linda Bushel (synthetic) | 1 Test Rd | Customer-Facing Estimate | 2026-10-27"
    assert AM.styles_sha256(out) == AM.styles_sha256(data)  # typography and styles are the master's
    # Fewer rows than the master leaves no placeholder behind; a missing token or table stops the document.
    out = AM.populate(AM.CUSTOMER_FACING_ESTIMATE, directory=tmp_path, tokens={"AMOUNT": "$1"},
                      tables={("Included Scope", "Description"): [("Only", "row")]})
    assert AM.PLACEHOLDER not in "".join(c.text for t in Document(io.BytesIO(out)).tables for r in t.rows for c in r.cells)
    with pytest.raises(AM.ApprovedMasterError) as raised:
        AM.populate(AM.CUSTOMER_FACING_ESTIMATE, directory=tmp_path, tokens={"NOT_IN_MASTER": "x"}, tables={})
    assert raised.value.block_code == AM.BLOCK_TOKEN_NOT_IN_MASTER
    with pytest.raises(AM.ApprovedMasterError) as raised:
        AM.populate(AM.CUSTOMER_FACING_ESTIMATE, directory=tmp_path, tokens={}, tables={("Invented", "Layout"): [("x", "y")]})
    assert raised.value.block_code == AM.BLOCK_TABLE_NOT_IN_MASTER


def test_the_real_approved_masters_01_02_03_populate_with_their_own_styles_header_and_logo():
    """Needs the confidential masters (office Mac default folder, or CALIBRAYTAI_APPROVED_MASTER_DIR)."""
    directory = AM.master_dir()
    for document_type, tables in (
        (AM.LABOUR_CALCULATION_DETAIL, {
            ("Labour Task", "Basis / Scope", "Hours", "Direct labour", "GM reference sell"): [("Framing", "deck", "8", "$600.00", "$705.88")],
            ("Metric", "Amount"): [("Total hours", "8")]}),
        (AM.INTERNAL_DETAILED_COST_BREAKDOWN, {
            ("Material / Package", "Direct Cost", "Status"): [("Lumber", "$100.00", "Written quote")],
            ("Task", "Hours", "Direct Cost"): [("Framing", "8", "$600.00")],
            ("Component", "Amount"): [("Total direct cost", "$700.00")]}),
        (AM.CUSTOMER_FACING_ESTIMATE, {
            ("Included Scope", "Description"): [("Deck", "Synthetic")],
            ("Price Component", "Amount"): [("Estimate subtotal", "$100.00")]}),
    ):
        source = AM.load_master_copy(document_type, directory)
        tokens = {"CLIENT": "Synthetic Client", "PROJECT": "Synthetic Project", "SITE": "1 Test Rd", "DATE": "2026-10-27"}
        out = AM.populate(document_type, directory=directory, tokens=tokens, tables=tables)
        assert AM.styles_sha256(out) == AM.styles_sha256(source)
        with zipfile.ZipFile(io.BytesIO(source)) as a, zipfile.ZipFile(io.BytesIO(out)) as b:
            media = sorted(n for n in a.namelist() if n.startswith("word/media/"))
            assert media and media == sorted(n for n in b.namelist() if n.startswith("word/media/"))
            for name in media + [n for n in a.namelist() if n.startswith("word/header")]:
                assert a.read(name) == b.read(name), name  # logo and header untouched


# Every PDF renderer in the office application, classified. A governed document type is produced from its
# master; the four renderers below still produce governed types and are listed for re-routing in
# docs/governance/approved-document-generation.md. An unclassified new renderer fails this test, so a new
# layout for a governed type cannot appear silently.
PENDING_REROUTE = {
    "proposal_pdf.py": AM.CLIENT_CONSTRUCTION_PROPOSAL,
    "guided_internal_cost_pdf.py": AM.INTERNAL_DETAILED_COST_BREAKDOWN,
    "estimate_quickbooks_pdf.py": AM.QUICKBOOKS_ESTIMATE_ENTRY,
    "brayman_v1_contract_presentation.py": AM.ONTARIO_CONSTRUCTION_CONTRACT,
}
NOT_GOVERNED = {
    "business_owner_assessment_pdf.py": "employment vs entrepreneurship calculator report",
    "desk_print.py": "Supplier Pro Desk print views",
    "permit_report_pdf.py": "permit context report",
    "signing_executed_pdf.py": "signature certificate added to an executed document",
    "supplier_package_pdf.py": "supplier order package",
    "supplier_quote_pdf.py": "supplier's quotation to a contractor",
}


def test_every_office_pdf_renderer_is_classified_and_governed_types_are_listed_for_re_routing():
    doc = (ROOT / "docs/governance/approved-document-generation.md").read_text(encoding="utf-8")
    found = {path.name for path in (ROOT / "app/services").glob("*.py")
             if re.search(r"SimpleDocTemplate|canvas\.Canvas|weasyprint", path.read_text(encoding="utf-8"))}
    assert found == set(PENDING_REROUTE) | set(NOT_GOVERNED), sorted(found ^ (set(PENDING_REROUTE) | set(NOT_GOVERNED)))
    assert not set(PENDING_REROUTE) & set(NOT_GOVERNED)
    for name, document_type in PENDING_REROUTE.items():
        assert f"`app/services/{name}`" in doc and document_type in AM.GOVERNED_DOCUMENT_TYPES
