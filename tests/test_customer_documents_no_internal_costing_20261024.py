"""Joel's decision (24 Oct 2026): internal labour rates and internal costing notes never reach a customer document.

A synthetic estimate carries a line with an internal labour-rate note, a line with an internal cost note, the
platform's own labour-rate note, a guided allowance-approval note and one note written for the customer. Each
customer-facing document is checked, HTML and PDF text: the office proposal preview and PDF, the customer's
proposal link and its PDF, the generated contract, and the change order PDF (which is also what a change order sends
for signature). The internal notes stay on the estimate lines (the internal costing record). Synthetic data only.
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from pypdf import PdfReader

from app import db
from app.models import CostItem
from app.models.estimate import EstimateLineItem, EstimateSection
from app.models.proposal import customer_line_note, is_internal_costing_note
from app.services import proposal_acceptance as PA
from app.services.estimate_builder import add_cost_item_line, add_manual_line
from app.services.organizations import DEFAULT_ORGANIZATION_ID as ORG
from tests.auth_fixtures import ensure_office_user, login_office_user
from tests.proposal_issue_support import approve_for_issue
from tests.test_proposal_client_acceptance import CLIENT, _version, app, client  # noqa: F401

RATE_NOTE = "Crew labour rate $85.00/hr, internal"
COST_NOTE = "Our cost from Valley Lumber $412.00 plus 18% markup"
PLATFORM_RATE_NOTE = "Company labour rate 85.00 per hour (standard 3, version 1). Hours entered by Ben Synthetic."
ALLOWANCE_NOTE = "Allowance approved by Ben Synthetic. Reason: The HVAC quote is still to come"
TYPED_NOTE = "Includes two coats of stain on the railing."
INTERNAL = (RATE_NOTE, COST_NOTE, PLATFORM_RATE_NOTE, ALLOWANCE_NOTE)
# Fragments that must not appear on any customer document (PDF text can wrap, so short and distinctive).
LEAKS = ("Valley Lumber", "85.00/hr", "labour rate", "Labour rate", "18% markup", "Hours entered by",
         "Allowance approved by", "HVAC quote is still to come")


def _pdf_text(data: bytes) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _clean(text, where):
    for leak in LEAKS:
        assert leak not in text, (where, leak)


def _estimate():
    project, version = _version("EST-INTERNAL-1")
    item = CostItem(organization_id=ORG, code="INT-DECK-BOARD", name="Deck boards (synthetic)", category="Material",
                    unit="ea", unit_cost=Decimal("12.00"), default_markup_percent=Decimal("25"))
    db.session.add(item)
    section = EstimateSection(estimate_version_id=version.id, name="Deck", sort_order=1, subtotal=Decimal("0"))
    db.session.add(section)
    db.session.commit()
    boards = add_cost_item_line(section, cost_item_id=item.id, quantity=40, notes=COST_NOTE, organization_id=ORG)
    railing = add_cost_item_line(section, cost_item_id=item.id, quantity=10, notes=TYPED_NOTE, organization_id=ORG)
    railing.description = "Railing"
    add_manual_line(section, line_type="Custom", description="Labour — build the deck", quantity=10, unit="hr",
                    unit_cost=85, notes=RATE_NOTE)
    add_manual_line(section, line_type="Custom", description="Labour — stairs", quantity=4, unit="hr", unit_cost=85,
                    notes=PLATFORM_RATE_NOTE)
    add_manual_line(section, line_type="Allowance", description="Lighting allowance", quantity=1, unit="allowance",
                    unit_cost=600, notes=ALLOWANCE_NOTE)
    db.session.commit()
    approve_for_issue(version)
    return project, version


def test_the_rule_keeps_customer_wording_and_drops_internal_costing():
    for note in INTERNAL:
        assert is_internal_costing_note(note) and customer_line_note(note) is None, note
    for note in ("Labour at $85 per hour", "Margin 22%", "Supplier price $4.10/ft2", "Overhead included",
                 "Wholesale from BMR"):
        assert customer_line_note(note) is None, note
    for note in (TYPED_NOTE, "Rule of thumb; quantity as discussed on site", "Pressure-treated, ground contact",
                 "Scope: Electrical: supply and install"):
        assert customer_line_note(note) == note


def test_no_customer_document_shows_an_internal_rate_or_costing_note(app, client):  # noqa: F811
    user = ensure_office_user()
    db.session.commit()
    login_office_user(client)
    project, version = _estimate()

    proposal = PA.sign_off_estimate(version, user=user)
    stored = [line.notes for section in proposal.sections for line in section.line_items]
    assert not set(INTERNAL) & set(stored) and TYPED_NOTE in stored

    # The office's view of the proposal (what will be sent) and its PDF.
    preview = client.get(f"/proposals/{proposal.id}/preview").get_data(as_text=True)
    office_pdf = _pdf_text(client.get(f"/proposals/{proposal.id}/pdf").data)
    for where, text in (("proposal preview", preview), ("proposal PDF", office_pdf)):
        _clean(text, where)
        assert TYPED_NOTE in text.replace("\n", " "), where

    # The customer's proposal link and its PDF.
    sent = PA.send_proposal_to_client(proposal, user=user, client_name=CLIENT[0], client_email=CLIENT[1])
    public = app.test_client()
    page = public.get(sent.path).get_data(as_text=True)
    link_pdf = _pdf_text(public.get(sent.path + "/pdf").data)
    for where, text in (("customer proposal link", page), ("customer proposal PDF", link_pdf)):
        _clean(text, where)
        assert TYPED_NOTE in text.replace("\n", " "), where

    # Accepted: the generated contract.
    accepted = public.post(sent.path + "/accept", data={"accept": "yes", "confirmed_name": CLIENT[0]})
    assert accepted.status_code == 200
    acceptance = PA.acceptance_for(db.session.get(PA.Proposal, proposal.id))
    if acceptance.contract_id is not None:
        from app.models.project_contract import ProjectContractSnapshot
        from app.services.contract_generation import retrieve_generated_contract_docx

        snapshot = ProjectContractSnapshot.query.filter_by(generated_contract_id=acceptance.contract_id).one()
        data = retrieve_generated_contract_docx(snapshot)  # the interim Brayman V1 contract is retained as a PDF
        _clean(_pdf_text(data) if data.startswith(b"%PDF") else _docx_text(data), "contract")
    assert acceptance.contract_id is not None, acceptance.contract_block_code

    # A change order made from the estimate, with an internal note typed on it.
    from app.project_controls.pdf import generate_change_order_pdf
    from app.project_controls.services import create_change_order

    version = db.session.get(type(version), version.id)
    co = create_change_order(project=project, title="Add a second stair (synthetic)", estimate_version=version,
                             copy_estimate_lines=True, notes=COST_NOTE, markup_percent=15, tax_percent=13,
                             organization_id=ORG)
    co_text = _pdf_text(generate_change_order_pdf(co).getvalue())
    _clean(co_text, "change order PDF")
    assert "Markup" not in co_text and "$85.00" not in co_text and "Unit Price" not in co_text
    assert "Labour — build the deck" in co_text and "Change order price before tax" in co_text

    # The internal costing record keeps every note.
    kept = {line.notes for line in EstimateLineItem.query.filter(
        EstimateLineItem.estimate_section_id.in_([s.id for s in version.sections])).all()}
    assert set(INTERNAL) <= kept


def test_a_change_order_with_selling_prices_and_no_markup_still_lists_them(app, client):  # noqa: F811
    from app.project_controls.pdf import generate_change_order_pdf
    from app.project_controls.services import add_change_order_item, create_change_order

    project, _version_row = _version("EST-INTERNAL-2")
    co = create_change_order(project=project, title="Extra post (synthetic)", tax_percent=13,
                             notes="Includes removal of the old post.", organization_id=ORG)
    add_change_order_item(co, description="Extra 6x6 post", quantity=1, unit="ea", unit_price="240.00")
    text = _pdf_text(generate_change_order_pdf(co).getvalue())
    assert "Unit Price" in text and "$240.00" in text and "Includes removal of the old post." in text
    assert "Markup" not in text
