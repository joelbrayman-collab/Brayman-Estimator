"""Family 05 DOCX merge from a frozen generated-contract snapshot.

Consumes frozen commercial facts and frozen legal-object bodies only.
Does not select jurisdiction, activate packages, or mutate source records.
Works from a byte copy of the governed master. Never writes the master path.
"""

from __future__ import annotations

import inspect
from dataclasses import dataclass
from io import BytesIO
from typing import Iterable, Optional

from docx import Document
from docx.text.paragraph import Paragraph

from app.services.family_05_master import (
    BLOCK_PRESENTATION_MASTER_SHA_MISMATCH,
    FAMILY_05_MASTER_SHA256,
    Family05MasterError,
    load_family_05_master_copy,
    sha256_bytes,
)

TOKEN_CLIENT = "CLIENT"
TOKEN_PROJECT = "PROJECT"
TOKEN_SITE = "SITE"
TOKEN_DATE = "DATE"
REQUIRED_COMMERCIAL_TOKENS = (TOKEN_CLIENT, TOKEN_PROJECT, TOKEN_SITE, TOKEN_DATE)

LEGAL_TEMPLATE_SLOT = (
    "Insert approved Ontario construction-contract template/version "
    "and obtain human approval."
)
WARRANTY_ATTACHMENT_SLOT = (
    "Attach only the separately approved warranty schedule."
)

SAFETY_LABELS = (
    "GOVERNED COMMERCIAL DRAFT",
    "COMMERCIAL DRAFT — NOT FOR EXECUTION",
    "NOT FOR SIGNATURE",
)

BLOCK_MISSING_COMMERCIAL_FACTS = "MISSING_COMMERCIAL_FACTS"
BLOCK_MISSING_REQUIRED_LEGAL_OBJECT = "MISSING_REQUIRED_LEGAL_OBJECT"
BLOCK_FAMILY_05_SLOT_UNRESOLVED = "FAMILY_05_SLOT_UNRESOLVED"
BLOCK_SAFETY_LABELS_STRIPPED = "FAMILY_05_SAFETY_LABELS_STRIPPED"

KIND_CONTRACT_PROVISION = "contract_provision"
KIND_WARRANTY = "warranty"


@dataclass(frozen=True)
class Family05MergeResult:
    merged: bool
    block_code: Optional[str]
    docx_bytes: Optional[bytes]
    artifact_sha256: Optional[str]


def _block(code: str) -> Family05MergeResult:
    return Family05MergeResult(
        merged=False,
        block_code=code,
        docx_bytes=None,
        artifact_sha256=None,
    )


def _ok(data: bytes) -> Family05MergeResult:
    return Family05MergeResult(
        merged=True,
        block_code=None,
        docx_bytes=data,
        artifact_sha256=sha256_bytes(data),
    )


def _paragraphs_in(container) -> Iterable[Paragraph]:
    for paragraph in getattr(container, "paragraphs", []) or []:
        yield paragraph
    for table in getattr(container, "tables", []) or []:
        for row in table.rows:
            for cell in row.cells:
                yield from _paragraphs_in(cell)


def _iter_all_paragraphs(document: Document) -> Iterable[Paragraph]:
    yield from _paragraphs_in(document)
    for section in document.sections:
        yield from _paragraphs_in(section.header)
        yield from _paragraphs_in(section.footer)


def _replace_token_in_paragraph(paragraph: Paragraph, token: str, value: str) -> None:
    if token not in (paragraph.text or ""):
        return
    runs = paragraph.runs
    if not runs:
        paragraph.add_run((paragraph.text or "").replace(token, value))
        return
    if len(runs) == 1:
        runs[0].text = (runs[0].text or "").replace(token, value)
        return
    joined = "".join(run.text or "" for run in runs)
    if token not in joined:
        return
    replaced = joined.replace(token, value)
    runs[0].text = replaced
    for run in runs[1:]:
        run.text = ""


def _set_paragraph_body(paragraph: Paragraph, body: str) -> None:
    lines = body.splitlines() or [""]
    if paragraph.runs:
        paragraph.runs[0].text = lines[0]
        for run in paragraph.runs[1:]:
            run.text = ""
        run = paragraph.runs[0]
    else:
        run = paragraph.add_run(lines[0])
        lines = lines[1:]
        for line in lines:
            run.add_break()
            run.add_text(line)
        return
    for line in lines[1:]:
        run.add_break()
        run.add_text(line)


def _replace_exact_slot(document: Document, placeholder: str, body: str) -> bool:
    found = False
    for paragraph in _iter_all_paragraphs(document):
        if (paragraph.text or "").strip() == placeholder:
            _set_paragraph_body(paragraph, body)
            found = True
            break
    return found


def _apply_commercial_tokens(document: Document, mapping: dict) -> None:
    for paragraph in _iter_all_paragraphs(document):
        for token, value in mapping.items():
            _replace_token_in_paragraph(paragraph, token, value)


def _document_plain_text(document: Document) -> str:
    return "\n".join(p.text or "" for p in _iter_all_paragraphs(document))


def _placeholder_still_present(rendered: str, token: str, value: str) -> bool:
    if token not in rendered:
        return False
    if token in value:
        return token in rendered.replace(value, "")
    return token in rendered


def _frozen_kind_body(legal_objects: Iterable[dict], kind: str) -> Optional[str]:
    for item in legal_objects:
        if (item.get("kind") or "") == kind:
            body = (item.get("body") or "").strip()
            if body:
                return body
    return None


def merge_family_05_from_frozen(
    *,
    presentation_master: dict,
    commercial: dict,
    legal_objects: Iterable[dict],
) -> Family05MergeResult:
    """Merge a copy of Family 05 using frozen snapshot values only."""
    expected = str((presentation_master or {}).get("sha256") or "").strip().lower()
    if expected != FAMILY_05_MASTER_SHA256:
        return _block(BLOCK_PRESENTATION_MASTER_SHA_MISMATCH)

    mapping = {
        TOKEN_CLIENT: str((commercial or {}).get("client_name") or "").strip(),
        TOKEN_PROJECT: str((commercial or {}).get("project_name") or "").strip(),
        TOKEN_SITE: str((commercial or {}).get("site") or commercial.get("project_address") or "").strip(),
        TOKEN_DATE: str((commercial or {}).get("contract_date") or "").strip(),
    }
    if any(not mapping[token] for token in REQUIRED_COMMERCIAL_TOKENS):
        return _block(BLOCK_MISSING_COMMERCIAL_FACTS)

    provision = _frozen_kind_body(legal_objects, KIND_CONTRACT_PROVISION)
    warranty = _frozen_kind_body(legal_objects, KIND_WARRANTY)
    if not provision or not warranty:
        return _block(BLOCK_MISSING_REQUIRED_LEGAL_OBJECT)

    try:
        master_bytes = load_family_05_master_copy(expected_sha256=expected)
    except Family05MasterError as exc:
        return _block(exc.block_code)

    document = Document(BytesIO(master_bytes))
    _apply_commercial_tokens(document, mapping)
    if not _replace_exact_slot(document, LEGAL_TEMPLATE_SLOT, provision):
        return _block(BLOCK_FAMILY_05_SLOT_UNRESOLVED)
    if not _replace_exact_slot(document, WARRANTY_ATTACHMENT_SLOT, warranty):
        return _block(BLOCK_FAMILY_05_SLOT_UNRESOLVED)

    rendered = _document_plain_text(document)
    for label in SAFETY_LABELS:
        if label not in rendered:
            return _block(BLOCK_SAFETY_LABELS_STRIPPED)
    leftover = [
        token
        for token in REQUIRED_COMMERCIAL_TOKENS
        if _placeholder_still_present(rendered, token, mapping[token])
    ]
    if leftover:
        return _block(BLOCK_FAMILY_05_SLOT_UNRESOLVED)

    out = BytesIO()
    document.save(out)
    return _ok(out.getvalue())


def family_05_merge_source() -> str:
    return inspect.getsource(merge_family_05_from_frozen)
